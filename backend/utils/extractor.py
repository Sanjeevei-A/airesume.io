"""
backend/utils/extractor.py
==========================
Extracts raw text from PDF, DOCX, and TXT resume files.
"""

import os
import fitz          # PyMuPDF
import pdfplumber
from docx import Document


# ─────────────────────────────────────────────────────────────────────────────
# Public API
# ─────────────────────────────────────────────────────────────────────────────

def extract_text(file_path: str) -> str:
    """
    Dispatch to the correct extractor based on file extension.
    Returns the full text as a single string.
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext == '.pdf':
        return _extract_pdf(file_path)
    elif ext == '.docx':
        return _extract_docx(file_path)
    elif ext == '.txt':
        return _extract_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


# ─────────────────────────────────────────────────────────────────────────────
# Private helpers
# ─────────────────────────────────────────────────────────────────────────────

def _extract_pdf(path: str) -> str:
    """
    Primary: PyMuPDF (fast).
    Fallback: pdfplumber (better for tables / complex layouts).
    """
    text = ""
    try:
        doc = fitz.open(path)
        for page in doc:
            text += page.get_text()
        doc.close()
    except Exception:
        text = ""

    # Fallback when PyMuPDF returns almost nothing
    if len(text.strip()) < 50:
        try:
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    raw = page.extract_text()
                    if raw:
                        text += raw + "\n"
        except Exception as e:
            raise RuntimeError(f"PDF extraction failed: {e}")

    return text.strip()


def _extract_docx(path: str) -> str:
    """Extract text from a Word document."""
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also grab text inside tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n".join(paragraphs)


def _extract_txt(path: str) -> str:
    """Read a plain-text file."""
    with open(path, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read()
